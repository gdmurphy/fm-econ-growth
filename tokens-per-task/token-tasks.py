import openpyxl
from collections import defaultdict

wb = openpyxl.load_workbook("tokens-per-task/Cost and Token count for GPT and Claude (1).xlsx", data_only=True)
ws = wb["Sheet1"]

headers = [cell.value for cell in ws[1]]
col = {h: i for i, h in enumerate(headers)}

# Aggregate across trials
model_data = defaultdict(lambda: {"input_tokens": 0, "output_tokens": 0, "task_count": 0, "n": 0})

for row in ws.iter_rows(min_row=2, values_only=True):
    model = row[col["Model"]]
    input_tok = row[col["Input (Tokens)"]]
    output_tok = row[col["Output (Tokens)"]]
    task_count = row[col["Task_Count"]]
    if model and input_tok and output_tok and task_count:
        model_data[model]["input_tokens"] += input_tok
        model_data[model]["output_tokens"] += output_tok
        model_data[model]["task_count"] += task_count
        model_data[model]["n"] += 1

print(f"{'Model':<35} {'Input/Output Ratio':>18}  {'Input Tok/Task':>14}  {'Output Tok/Task':>15}  {'Total Tok/Task':>14}")
print("-" * 105)

all_ratios = []
all_input_per_task = []
all_output_per_task = []
all_total_per_task = []

for model, d in model_data.items():
    inp = d["input_tokens"]
    out = d["output_tokens"]
    tasks = d["task_count"]

    ratio = inp / out
    input_per_task = inp / tasks
    output_per_task = out / tasks
    total_per_task = (inp + out) / tasks

    all_ratios.append(ratio)
    all_input_per_task.append(input_per_task)
    all_output_per_task.append(output_per_task)
    all_total_per_task.append(total_per_task)

    print(f"{model:<35} {ratio:>18.4f}  {input_per_task:>14.1f}  {output_per_task:>15.1f}  {total_per_task:>14.1f}")

n = len(all_ratios)
print("-" * 105)
print(f"{'AVERAGE':<35} {sum(all_ratios)/n:>18.4f}  {sum(all_input_per_task)/n:>14.1f}  {sum(all_output_per_task)/n:>15.1f}  {sum(all_total_per_task)/n:>14.1f}")
